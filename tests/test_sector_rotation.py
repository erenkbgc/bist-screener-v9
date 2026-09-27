"""RRG: kadran mantigi, future_star tespiti, yetersiz veri."""
import numpy as np
import pandas as pd

from core.sector_rotation import compute_rrg, quadrant, sector_index_for, annotate_candidates


def _series(values):
    idx = pd.bdate_range("2023-01-02", periods=len(values))
    return pd.Series(values, index=idx, dtype=float)


def test_quadrants():
    assert quadrant(101, 101) == "Leading"
    assert quadrant(101, 99) == "Weakening"
    assert quadrant(99, 99) == "Lagging"
    assert quadrant(99, 101) == "Improving"


def test_insufficient_history_returns_none():
    assert compute_rrg(_series(np.linspace(100, 110, 100)), _series(np.full(100, 100.0))) is None


def test_turnaround_sector_rotates_clockwise_through_improving():
    # 2 yil boyunca endeksi geride birakip dusen, sonra donen sektor. RRG
    # saat yonunde doner: Lagging -> Improving -> Leading.
    n = 520
    bench = _series(np.full(n, 100.0))
    rel = np.concatenate([np.linspace(1.0, 0.70, n - 20), np.linspace(0.70, 0.74, 20)])
    rrg = compute_rrg(_series(100.0 * rel), bench)
    assert rrg is not None
    quads = [quadrant(r, m) for r, m in rrg["tail"]]
    assert "Improving" in quads
    assert quads.index("Improving") < len(quads) - 1
    assert rrg["prior_quadrant"] == "Improving"
    # dusus surerken (donusten once) goreli guc 100 altinda
    early = compute_rrg(_series(100.0 * rel[: n - 20]), bench.iloc[: n - 20])
    assert early["rs_ratio"] < 100


def test_sector_mapping_and_fallback():
    assert sector_index_for("Bankacılık", "XUHIZ") == ("XBANK", "sector")
    assert sector_index_for("BILINMIYOR", "XUSIN") == ("XUSIN", "supersector_fallback")
    assert sector_index_for(None, None) == (None, "unmapped")
    c = [{"sector": "Sigorta", "supersector": "XUHIZ"}]
    annotate_candidates(c, {"XSGRT": {"quadrant": "Improving", "future_star": True}})
    assert c[0]["rrg_quadrant"] == "Improving" and c[0]["rrg_future_star"] is True


def test_sector_excel_builds_with_chart_and_candidates(tmp_path):
    from openpyxl import load_workbook
    from report.sector_excel import write_sector_rotation_excel
    rot = {
        "XSGRT": {"week_end": "2026-09-25", "rs_ratio": 99.21, "rs_momentum": 100.91, "quadrant": "Improving",
                  "weeks_in_quadrant": 2, "prior_quadrant": "Lagging", "future_star": True,
                  "tail": [[98.5, 99.2], [98.7, 99.8], [98.9, 100.3], [99.1, 100.7], [99.21, 100.91]]},
        "XBANK": {"week_end": "2026-09-25", "rs_ratio": 101.32, "rs_momentum": 101.92, "quadrant": "Leading",
                  "weeks_in_quadrant": 2, "prior_quadrant": "Improving", "future_star": False,
                  "tail": [[99.8, 101.0], [100.2, 101.5], [100.7, 101.8], [101.0, 102.0], [101.32, 101.92]]},
    }
    cands = [{"ticker": "AKBNK", "sector_index": "XBANK", "candidate_state": "OPPORTUNITY"},
             {"ticker": "XXX", "sector_index": "XBANK", "candidate_state": "NO_ACTION"}]
    path = write_sector_rotation_excel(tmp_path / "r.xlsx", "2026-09-26", rot, cands)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Özet", "Kuyruk", "RRG Grafik", "Yöntem"]
    ws = wb["Özet"]
    assert ws["A6"].value == "XSGRT ★"          # Iyilesen ilk sirada
    assert ws["C6"].value == "İyileşen"
    assert ws["M7"].value == "AKBNK"             # NO_ACTION adaylar listelenmez
    assert len(wb["RRG Grafik"]._charts) == 1
    assert write_sector_rotation_excel(tmp_path / "e.xlsx", "2026-09-26", {}) is None
