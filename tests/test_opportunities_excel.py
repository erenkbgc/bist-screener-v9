"""Firsatlar Excel'i: payload'dan uretilir, durum sirasi ve bos gun korunur."""
from openpyxl import load_workbook

from report.opportunities_excel import write_opportunities_excel


def _c(ticker, state, score, bucket="long_term"):
    return {"ticker": ticker, "bucket": bucket, "candidate_state": state, "final_score": score,
            "entry_price": 10.0, "target_price": 11.0, "fair_value_base": 12.0, "sector": "S",
            "factor_attribution": {"top_positive_factor": "Değerleme"}}


def test_sheets_order_and_values(tmp_path):
    payload = {
        "passing_candidates": [_c("WWW", "WATCHLIST", 0.9), _c("SSS", "STRONG_OPPORTUNITY", 0.1),
                               _c("OOO", "OPPORTUNITY", 0.5), _c("KKK", "OPPORTUNITY", 0.2, "short_term")],
        "run_health": {"status": "DEGRADED", "issues": ["kap_unavailable"]},
        "regime": {"xu100_level": 13000.0, "bond_2y_pct": 40.0, "usdtry_spot": 49.0},
        "weights": {"valuation_z": 0.5, "catalyst_score": 0.25},
    }
    path = write_opportunities_excel(tmp_path / "f.xlsx", "2026-10-02", payload)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Özet", "Uzun Vade", "Kısa Vade", "Açıklamalar"]
    lt = wb["Uzun Vade"]
    # durum sirasi skordan once gelir: guclu > firsat > izleme
    assert lt["A1"].value.startswith("BİLGİ AMAÇLI")
    assert [lt.cell(row=r, column=1).value for r in (3, 4, 5)] == ["SSS", "OOO", "WWW"]
    headers = [c.value for c in lt[2]]
    gap = lt.cell(row=3, column=headers.index("Adil değer farkı %") + 1).value
    assert abs(gap - 20.0) < 1e-9
    assert wb["Kısa Vade"].cell(row=2, column=1).value == "KKK"
    summary = {r[0]: r[1] for r in wb["Özet"].iter_rows(min_row=3, values_only=True) if r[0]}
    assert summary["Veri sağlığı"].startswith("Kısmi")


def test_empty_day(tmp_path):
    wb = load_workbook(write_opportunities_excel(tmp_path / "e.xlsx", "2026-10-02", {}))
    assert "aday yok" in wb["Uzun Vade"].cell(row=3, column=1).value


def test_rotation_sheet_first_and_labelled(tmp_path):
    payload = {"rotation": {"n_holdings": 1, "buy": [{"ticker": "AAA", "valuation_z": 1.2, "rank_pct": 0.95}],
                            "hold": [], "sell": [], "rebalance_date": "2026-11-02"},
               "passing_candidates": [{**_c("AAA", "WATCHLIST", 0.3), "unknown_checks": ["valuation_excess_pct"]}]}
    wb = load_workbook(write_opportunities_excel(tmp_path / "r.xlsx", "2026-11-02", payload))
    assert wb.sheetnames[:2] == ["Özet", "Rotasyon"]
    assert wb["Rotasyon"]["A1"].value.startswith("TEST EDİLEN TEK KARAR KURALI")
    assert wb["Rotasyon"].cell(row=3, column=2).value == "AAA"
    lt = wb["Uzun Vade"]
    headers = [c.value for c in lt[2]]
    assert lt.cell(row=3, column=headers.index("Eksik veri") + 1).value == "valuation_excess_pct"


def test_single_stock_check_sheet(tmp_path):
    """plan_2026-10-21 P2-4: rotasyon hisseleri icin ayri, bilgi amacli kontrol sayfasi."""
    from core.rotation import holding_checks
    cands = [{"ticker": "AAA", "sector": "S", "loss_flag": True, "net_debt_ebitda": 4.2,
              "value_inputs": {"stmt_period": "2026Q2"}, "unknown_checks": [], "risk_flags": ["Zarar"]}]
    checks = holding_checks(cands, [{"ticker": "AAA"}, {"ticker": "ZZZ"}])
    assert checks[0]["statement_period"] == "2026Q2" and checks[0]["in_today_candidates"]
    assert checks[1]["in_today_candidates"] is False
    payload = {"rotation": {"n_holdings": 2, "buy": [{"ticker": "AAA", "valuation_z": 1.0, "rank_pct": 0.9}],
                            "hold": [], "sell": [], "checks": checks}}
    wb = load_workbook(write_opportunities_excel(tmp_path / "c.xlsx", "2026-11-02", payload))
    assert wb.sheetnames[:3] == ["Özet", "Rotasyon", "Tek Hisse Kontrol"]
    ws = wb["Tek Hisse Kontrol"]
    assert ws["A1"].value.startswith("BİLGİ AMAÇLI")
    headers = [c.value for c in ws[2]]
    assert ws.cell(row=3, column=headers.index("Net borç / FAVÖK") + 1).value == 4.2
    assert ws.cell(row=3, column=headers.index("Zarar") + 1).value == "Evet"
