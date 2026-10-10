"""Gunun firsatlari Excel eki: e-postayi acmadan tek dosyada okunur ozet.

Sayfalar:
  Ozet        : tarih, veri sagligi, rejim, aday sayilari, durum aciklamalari
  Uzun Vade   : 180 gunluk adaylar (durum + skora gore sirali)
  Kisa Vade   : 20 gunluk adaylar
  Aciklamalar : kolon tanimlari ve sinirlar
Veri kaynagi yalnizca payload (passing_candidates, run_health, regime):
e-postadaki sayilarla birebir aynidir.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

STATE_TR = {"STRONG_OPPORTUNITY": "Güçlü fırsat", "OPPORTUNITY": "Fırsat", "WATCHLIST": "İzleme"}
STATE_ORDER = {"STRONG_OPPORTUNITY": 0, "OPPORTUNITY": 1, "WATCHLIST": 2}
STATE_FILL = {"STRONG_OPPORTUNITY": "C6EFCE", "OPPORTUNITY": "E2EFDA", "WATCHLIST": "FFF2CC"}
QUADRANT_TR = {"Leading": "Lider", "Improving": "İyileşen", "Weakening": "Zayıflayan", "Lagging": "Geride"}
HEALTH_TR = {"OK": "Tamam", "DEGRADED": "Kısmi (bazı veri kaynakları eksik)",
             "FAILED": "YETERSİZ: sinyaller güvenilir değil"}
HEALTH_FILL = {"OK": "C6EFCE", "DEGRADED": "FFF2CC", "FAILED": "F8CBAD"}
ISSUE_TR = {
    "fundamentals_coverage": "Temel veri kapsamı düşük",
    "price_fresh_coverage": "Güncel fiyat kapsamı düşük",
    "macro_missing": "Makro veri eksik",
    "index_series_empty": "BIST 100 serisi boş",
    "data_filter_share": "Adayların önemli kısmı veri eksikliğinden elendi",
    "kap_unavailable": "KAP'a erişilemedi (rapor tarihleri)",
    "catalyst_redistributed": "Katalizör verisi eksik; ağırlığı diğer faktörlere dağıtıldı",
}

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

PCT = "0.0"      # yuzde puan (orn. 22.3 = %22.3)
PRICE = "#,##0.00"
SCORE = "0.000"


def _pct_gap(a, b):
    try:
        return (a / b - 1) * 100 if a and b else None
    except (TypeError, ZeroDivisionError):
        return None


def _bool_tr(v):
    return "Evet" if v else ("Hayır" if v is not None else None)


# (baslik, deger fonksiyonu, sayi formati, genislik)
LONG_COLS = [
    ("Hisse", lambda c: c.get("ticker"), None, 9),
    ("Sektör", lambda c: c.get("sector"), None, 22),
    ("Durum", lambda c: STATE_TR.get(c.get("candidate_state"), c.get("candidate_state")), None, 13),
    ("Skor", lambda c: c.get("final_score"), SCORE, 8),
    ("Fiyat", lambda c: c.get("entry_price"), PRICE, 10),
    ("Alış bandı alt", lambda c: c.get("entry_low"), PRICE, 11),
    ("Alış bandı üst", lambda c: c.get("entry_high"), PRICE, 11),
    ("Stop", lambda c: c.get("stop_loss"), PRICE, 10),
    ("Hedef (180g)", lambda c: c.get("target_price"), PRICE, 11),
    ("Beklenen getiri %", lambda c: c.get("expected_roi_pct"), PCT, 11),
    ("Hurdle %", lambda c: c.get("hurdle_rate_pct"), PCT, 9),
    ("Hurdle üstü %", lambda c: c.get("excess_over_hurdle_pct"), PCT, 10),
    ("Hedefe ulaşma olasılığı %", lambda c: c.get("target_hit_prob_pct"), PCT, 12),
    ("Adil değer (baz)", lambda c: c.get("fair_value_base"), PRICE, 11),
    ("Adil değer farkı %", lambda c: _pct_gap(c.get("fair_value_base"), c.get("entry_price")), PCT, 11),
    ("Adil değer düşük", lambda c: c.get("fair_value_low"), PRICE, 11),
    ("Adil değer yüksek", lambda c: c.get("fair_value_high"), PRICE, 11),
    ("Pozisyon %", lambda c: c.get("position_size_pct"), PCT, 9),
    ("Destek", lambda c: c.get("nearest_support"), PRICE, 9),
    ("Direnç", lambda c: c.get("nearest_resistance"), PRICE, 9),
    ("Direnç hedeften önce", lambda c: _bool_tr(c.get("resistance_before_target")), None, 10),
    ("Değer tuzağı riski", lambda c: _bool_tr(c.get("value_trap_risk")), None, 10),
    ("5 yıllık zirveden %", lambda c: c.get("drawdown_from_5y_peak_pct"), PCT, 10),
    ("52 haftalık zirveden %", lambda c: c.get("dist_52w_high_pct"), PCT, 10),
    ("Risk bayrakları", lambda c: ", ".join(c.get("risk_flags") or []) or None, None, 30),
    ("Sektör RRG", lambda c: QUADRANT_TR.get(c.get("rrg_quadrant"), c.get("rrg_quadrant")), None, 11),
    ("F/K", lambda c: c.get("pe"), "0.0", 7),
    ("PD/DD", lambda c: c.get("pb"), "0.00", 7),
    ("Piotroski (0-1)", lambda c: c.get("piotroski_normalized_score"), "0.00", 9),
    ("Beta", lambda c: c.get("beta_60_120d"), "0.00", 7),
    ("En güçlü faktör", lambda c: (c.get("factor_attribution") or {}).get("top_positive_factor"), None, 22),
    ("En zayıf faktör", lambda c: (c.get("factor_attribution") or {}).get("top_negative_factor"), None, 22),
]
SHORT_COLS = [col for col in LONG_COLS if col[0] not in (
    "Hedef (180g)", "Adil değer (baz)", "Adil değer farkı %", "Adil değer düşük", "Adil değer yüksek",
    "Değer tuzağı riski", "5 yıllık zirveden %", "52 haftalık zirveden %", "Risk bayrakları",
    "F/K", "PD/DD", "Piotroski (0-1)")]
SHORT_COLS.insert(8, ("Hedef (20g)", lambda c: c.get("target_price"), PRICE, 11))
SHORT_COLS.insert(9, ("Hacim oranı (20g)", lambda c: c.get("volume_ratio_20d"), "0.00", 10))


def _header(ws, row: int, cols) -> None:
    for j, (title, _, _, width) in enumerate(cols, 1):
        cell = ws.cell(row=row, column=j, value=title)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(j)].width = width
    ws.row_dimensions[row].height = 32


def _candidate_sheet(wb: Workbook, title: str, cands: list[dict], cols) -> None:
    ws = wb.create_sheet(title)
    _header(ws, 1, cols)
    cands = sorted(cands, key=lambda c: (STATE_ORDER.get(c.get("candidate_state"), 9),
                                         -(c.get("final_score") or 0)))
    for i, c in enumerate(cands, 2):
        fill = PatternFill("solid", fgColor=STATE_FILL.get(c.get("candidate_state"), "FFFFFF"))
        for j, (_, fn, fmt, _) in enumerate(cols, 1):
            try:
                v = fn(c)
            except Exception:  # noqa: BLE001 - tek hucre hatasi tabloyu bozmasin
                v = None
            cell = ws.cell(row=i, column=j, value=v)
            cell.border = BORDER
            if fmt and isinstance(v, (int, float)) and not isinstance(v, bool):
                cell.number_format = fmt
            if j <= 3:
                cell.fill = fill
    if not cands:
        ws.cell(row=2, column=1, value="Bugün bu vadede kapılardan geçen aday yok.")
    ws.freeze_panes = "B2"
    if cands:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(cands) + 1}"


def build_opportunities_workbook(as_of_date: str, payload: dict) -> Workbook:
    cands = payload.get("passing_candidates") or []
    long_c = [c for c in cands if c.get("bucket") == "long_term"]
    short_c = [c for c in cands if c.get("bucket") == "short_term"]
    health = payload.get("run_health") or {}
    regime = payload.get("regime") or {}
    if isinstance(regime, list):
        regime = regime[0] if regime else {}

    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    ws["A1"] = f"BIST Günlük Fırsatlar — {as_of_date}"
    ws["A1"].font = Font(bold=True, size=14)
    rows = [("Veri sağlığı", HEALTH_TR.get(health.get("status"), "Ölçülmedi"))]
    for issue in health.get("issues") or []:
        rows.append(("  Sorun", ISSUE_TR.get(issue, issue)))
    rows += [
        ("BIST 100", regime.get("xu100_level")),
        ("2Y tahvil %", regime.get("bond_2y_pct")),
        ("USD/TRY", regime.get("usdtry_spot")),
        ("Uzun vade aday", len(long_c)),
        ("Kısa vade aday", len(short_c)),
    ]
    for state, label in STATE_TR.items():
        rows.append((f"  {label}", sum(1 for c in cands if c.get("candidate_state") == state)))
    r = 3
    for k, v in rows:
        ws.cell(row=r, column=1, value=k).font = Font(bold=not k.startswith("  "))
        cell = ws.cell(row=r, column=2, value=v)
        if k == "Veri sağlığı" and health.get("status") in HEALTH_FILL:
            cell.fill = PatternFill("solid", fgColor=HEALTH_FILL[health["status"]])
            cell.font = Font(bold=True)
        r += 1
    r += 1
    ws.cell(row=r, column=1, value="Durumlar").font = Font(bold=True)
    for state, text in (
        ("STRONG_OPPORTUNITY", "Yüksek güven, kovasında ilk %25, Piotroski güçlü, değer tuzağı işareti yok"),
        ("OPPORTUNITY", "Kapılardan geçti, hurdle üstü"),
        ("WATCHLIST", "Kapılardan geçti ama öne çıkmıyor; izle"),
    ):
        r += 1
        ws.cell(row=r, column=1, value=STATE_TR[state]).fill = PatternFill("solid", fgColor=STATE_FILL[state])
        ws.cell(row=r, column=2, value=text)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 90

    _candidate_sheet(wb, "Uzun Vade", long_c, LONG_COLS)
    _candidate_sheet(wb, "Kısa Vade", short_c, SHORT_COLS)
    if any(c.get("experimental") for c in short_c):
        wk = wb["Kısa Vade"]
        wk.insert_rows(1)
        wk["A1"] = ("DENEYSEL, FIRSAT DEĞİLDİR: kısa vade kuralı geçmiş veride piyasa ortalamasının gerisinde kaldı; "
                    "önceden kayıtlı alternatiflerin hiçbiri testi geçmedi. Yalnızca izleme.")
        wk["A1"].font = Font(bold=True, color="9C5700")
        wk.freeze_panes = "B3"
        if wk.auto_filter.ref:
            wk.auto_filter.ref = wk.auto_filter.ref.replace("A1:", "A2:")

    from core import diagnostics as dg
    from core.targets import convergence_alpha
    w = payload.get("weights") or {}
    alpha = convergence_alpha()
    wm = wb.create_sheet("Açıklamalar")
    notes = [
        "Yatırım tavsiyesi değildir. Model çıktısıdır; gerçekleşen performans henüz ölçülmedi (ilk sonuçlar 2026-10-01'den itibaren).",
        (f"Skor = {w.get('valuation_z', 0):g} değerleme + {w.get('catalyst_score', 0):g} katalizör + "
         f"{w.get('ownership_quality_z', 0):g} ortaklık + {w.get('low_vol_z', 0):g} düşük volatilite + "
         f"{w.get('momentum_z', 0):g} momentum. KAP verisi alınamayan hissede katalizör ağırlığı diğerlerine dağıtılır."),
        (f"Hedef (180g) = (fiyat + {alpha:g} × (adil değer − fiyat)) × (1 + özsermaye maliyeti)^(180/365) − temettü. "
         f"{alpha:g}, geçmiş verideki gerçek yakınsama hızından ölçüldü (fiyat 6 ayda adil değer farkının ~%3'ünü kapattı)."),
        f"Beklenen getirinin büyük kısmı piyasa getirisi beklentisidir (özsermaye maliyeti); hisseye özgü kısım 'Adil değer farkı × {alpha:g}' kadardır.",
        "Hurdle % = 180 günlük risksiz getiri (2Y tahvil). Hurdle üstü % = beklenen getiri − hurdle.",
        "Hedefe ulaşma olasılığı modelden gelir (lognormal); gerçekleşmiş isabet oranı değildir.",
        ("Alış bandı ve stop ATR tabanlıdır; pozisyon % sabit risk kuralından (hesabın ~%1.5'i risk) türetilir. "
         "Alış bandı yalnızca bilgidir: on-kayıtlı test (docs/research/entry_timing_preregistration.md) "
         "geri çekilme/limit, RSI, trend onayı ve kademeli alım kurallarının hemen almaktan iyi olmadığını gösterdi."),
        "Değer tuzağı riski: fiyat trendi aşağı; ucuzluk düşüşün sonucu olabilir.",
        (f"Zirveden % ve risk bayrakları yalnızca bilgidir, skoru değiştirmez: 5 yıllık zirveden "
         f"%{-dg.CRASH_FROM_PEAK_PCT:g}+ düşüş, 52 haftalık zirveden %{-dg.FAR_FROM_52W_HIGH_PCT:g}+ uzaklık, "
         f"zarar (TTM HBK < 0), borç/sıkıntı (negatif özsermaye, negatif FAVÖK + net borç, "
         f"net borç/FAVÖK > {dg.NET_DEBT_EBITDA_MAX:g} veya FAVÖK/finansman gideri < {dg.INTEREST_COVERAGE_MIN:g}). "
         "Eşikler henüz test edilmedi."),
        "Sektör RRG: sektörün BIST 100'e göre göreli gücü; skora dahil değil, tek başına alım sinyali değil.",
        "Backtest kanıtı (2013–2026): değer ve düşük volatilite faktörleri kesitsel olarak çalışıyor; hayatta kalma yanlılığı ve nominal TL uyarıları geçerli.",
    ]
    wm["A1"] = "Kolonlar ve sınırlar"
    wm["A1"].font = Font(bold=True, size=12)
    for i, n in enumerate(notes, 3):
        c = wm.cell(row=i, column=1, value=f"• {n}")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    wm.column_dimensions["A"].width = 130
    return wb


def write_opportunities_excel(path: Path, as_of_date: str, payload: dict) -> Path:
    wb = build_opportunities_workbook(as_of_date, payload)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
