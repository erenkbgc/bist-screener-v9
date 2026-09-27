"""Sektor rotasyonu Excel eki (yatirimci sunumu).

Sayfalar:
  Ozet        : kadran renkli tablo, yon, future star, sektordeki adaylar
  RRG Grafik  : 5 haftalik kuyruklu dagilim grafigi (sektor nereye kayiyor)
  Kuyruk      : grafigin veri tablosu
  Yontem      : formul ve sinirlar
Veri kaynagi yalnizca payload["sector_rotation"] ve passing_candidates --
e-postadaki sayilarla birebir aynidir.
"""
from __future__ import annotations

import math
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.marker import Marker
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SECTOR_NAMES_TR = {
    "XBANK": "Banka", "XHOLD": "Holding ve Yatırım", "XGIDA": "Gıda, İçecek",
    "XKMYA": "Kimya, Petrol, Plastik", "XMANA": "Metal Ana", "XMESY": "Metal Eşya, Makine",
    "XTAST": "Taş, Toprak", "XTEKS": "Tekstil, Deri", "XKAGT": "Orman, Kağıt, Basım",
    "XELKT": "Elektrik", "XINSA": "İnşaat", "XTRZM": "Turizm", "XULAS": "Ulaştırma",
    "XILTM": "İletişim", "XBLSM": "Bilişim", "XTCRT": "Ticaret", "XSGRT": "Sigorta",
    "XGMYO": "GYO", "XFINK": "Finansal Kiralama, Faktoring", "XAKUR": "Aracı Kurumlar",
    "XSPOR": "Spor", "XMADN": "Madencilik", "XYORT": "Menkul Kıymet Yat. Ort.",
    "XUSIN": "Sinai (genel)", "XUHIZ": "Hizmetler (genel)",
}
QUADRANT_ORDER = ["Improving", "Leading", "Weakening", "Lagging"]
QUADRANT_TR = {"Leading": "Lider", "Improving": "İyileşen", "Weakening": "Zayıflayan", "Lagging": "Geride"}
QUADRANT_FILL = {"Leading": "C6EFCE", "Improving": "DDEBF7", "Weakening": "FFF2CC", "Lagging": "F8CBAD"}
# Yorumlar backtest kanitina gore (scripts/backtest_sector_rotation.py,
# 2013-2026, 23 sektor endeksi): Iyilesen/erken ivme sepetleri esit agirlikli
# sektor sepetinin yilda ~%5-8 GERISINDE kaldi; 12-1 sektor momentumu (Lider
# tarafi) EW'yi ~%9-13 gecti ama coklu test esiginin altinda (t~1.5-1.9).
QUADRANT_ACTION = {
    "Improving": "İvme dönüyor; tek başına alım sinyali değil (backtest: EW gerisinde)",
    "Leading": "Güçlü göreli trend; momentum tarafı (backtest: EW üstü, anlamlı değil)",
    "Weakening": "İvme kayboluyor; stopları gözden geçir",
    "Lagging": "Zayıf göreli güç",
}
NEXT_QUADRANT = {"Lagging": "Improving", "Improving": "Leading", "Leading": "Weakening", "Weakening": "Lagging"}

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _direction(tail: list[list[float]]) -> tuple[str, float, float]:
    """Kuyrugun ilk->son noktasina gore yon: ok + (dR, dM)."""
    if not tail or len(tail) < 2:
        return "-", 0.0, 0.0
    dr = tail[-1][0] - tail[0][0]
    dm = tail[-1][1] - tail[0][1]
    angle = math.degrees(math.atan2(dm, dr))
    arrows = [(22.5, "→"), (67.5, "↗"), (112.5, "↑"), (157.5, "↖")]
    a = abs(angle)
    arrow = "←"
    for lim, sym in arrows:
        if a <= lim:
            arrow = sym
            break
    if angle < 0 and arrow in ("↗", "↑", "↖"):
        arrow = {"↗": "↘", "↑": "↓", "↖": "↙"}[arrow]
    return arrow, round(dr, 2), round(dm, 2)


def _header(ws, row: int, headers: list[str]) -> None:
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def build_sector_rotation_workbook(as_of_date: str, sector_rotation: dict[str, dict],
                                   candidates: list[dict] | None = None) -> Workbook:
    candidates = candidates or []
    by_sector: dict[str, list[str]] = {}
    for c in candidates:
        code = c.get("sector_index")
        if code and c.get("candidate_state") in ("STRONG_OPPORTUNITY", "OPPORTUNITY", "WATCHLIST"):
            by_sector.setdefault(code, [])
            if c["ticker"] not in by_sector[code]:
                by_sector[code].append(c["ticker"])

    rows = sorted(sector_rotation.items(),
                  key=lambda kv: (QUADRANT_ORDER.index(kv[1]["quadrant"]), -kv[1]["rs_momentum"]))

    wb = Workbook()
    ws = wb.active
    ws.title = "Özet"
    ws["A1"] = f"BIST Sektör Rotasyonu (RRG) — {as_of_date}"
    ws["A1"].font = Font(bold=True, size=16, color="1F3864")
    ws["A2"] = ("Benchmark XU100 · haftalık veri · 100 merkezli. ★ = son haftalarda Geride'den "
                "İyileşen'e geçen erken ivme sektörü. Bilgi amaçlıdır, yatırım tavsiyesi değildir.")
    ws["A2"].font = Font(italic=True, color="595959")
    ws.merge_cells("A2:L2")

    stars = [QUADRANT_TR.get(v["quadrant"]) and k for k, v in rows if v.get("future_star")]
    ws["A3"] = "Erken ivme (future star): " + (", ".join(
        f"{k} {SECTOR_NAMES_TR.get(k, '')}" for k in stars) if stars else "yok")
    ws["A3"].font = Font(bold=True, color="2F5597")

    headers = ["Endeks", "Sektör", "Kadran", "RS-Ratio", "RS-Momentum", "Kadranda (hafta)",
               "Önceki Kadran", "Yön (5h)", "ΔRS-Ratio (5h)", "ΔMomentum (5h)",
               "Olası Sonraki Kadran", "Yorum", "Rapordaki Adaylar"]
    _header(ws, 5, headers)
    for i, (code, v) in enumerate(rows, start=6):
        arrow, dr, dm = _direction(v.get("tail") or [])
        values = [
            code + (" ★" if v.get("future_star") else ""),
            SECTOR_NAMES_TR.get(code, code),
            QUADRANT_TR[v["quadrant"]],
            v["rs_ratio"], v["rs_momentum"], v["weeks_in_quadrant"],
            QUADRANT_TR.get(v.get("prior_quadrant") or "", "-"),
            arrow, dr, dm,
            QUADRANT_TR[NEXT_QUADRANT[v["quadrant"]]],
            QUADRANT_ACTION[v["quadrant"]],
            ", ".join(by_sector.get(code, [])) or "-",
        ]
        fill = PatternFill("solid", fgColor=QUADRANT_FILL[v["quadrant"]])
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=i, column=col, value=val)
            cell.border = BORDER
            cell.fill = fill
            if col in (4, 5, 9, 10):
                cell.number_format = "0.00"
            if col == 8:
                cell.alignment = Alignment(horizontal="center")
                cell.font = Font(size=14, bold=True)
        if v.get("future_star"):
            ws.cell(row=i, column=1).font = Font(bold=True, color="C00000")
    for col, width in enumerate([10, 26, 12, 10, 12, 10, 13, 9, 12, 13, 15, 38, 30], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A6"
    ws.row_dimensions[5].height = 32

    # --- Kuyruk verisi + grafik ---
    wt = wb.create_sheet("Kuyruk")
    wt["A1"] = "Son 5 haftalık RRG yolu (en eski → en yeni)"
    wt["A1"].font = Font(bold=True, size=12)
    _header(wt, 3, ["Endeks", "Hafta", "RS-Ratio", "RS-Momentum"])
    r = 4
    ranges: list[tuple[str, int, int]] = []
    for code, v in rows:
        tail = v.get("tail") or []
        start = r
        for k, (rr, mm) in enumerate(tail):
            wt.cell(row=r, column=1, value=code)
            wt.cell(row=r, column=2, value=k - len(tail) + 1)  # 0 = bu hafta
            wt.cell(row=r, column=3, value=rr).number_format = "0.00"
            wt.cell(row=r, column=4, value=mm).number_format = "0.00"
            r += 1
        if tail:
            ranges.append((code, start, r - 1))
    for col, width in enumerate([10, 8, 12, 14], 1):
        wt.column_dimensions[get_column_letter(col)].width = width

    wc = wb.create_sheet("RRG Grafik")
    wc["A1"] = "Göreli Rotasyon Grafiği — her çizgi son 5 hafta, çizgi sonu = bu hafta"
    wc["A1"].font = Font(bold=True, size=12)
    wc["A2"] = ("Sağ-üst: Lider · Sol-üst: İyileşen · Sağ-alt: Zayıflayan · Sol-alt: Geride. "
                "Sektörler genelde saat yönünde döner: Geride → İyileşen → Lider → Zayıflayan.")
    wc["A2"].font = Font(italic=True, color="595959")
    chart = ScatterChart()
    chart.title = f"BIST Sektör RRG — {as_of_date}"
    chart.style = 2
    chart.x_axis.title = "RS-Ratio (göreli güç)"
    chart.y_axis.title = "RS-Momentum (ivme)"
    all_r = [p[0] for _, v in rows for p in (v.get("tail") or [])] or [100]
    all_m = [p[1] for _, v in rows for p in (v.get("tail") or [])] or [100]
    span_r = max(abs(x - 100) for x in all_r) + 0.5
    span_m = max(abs(x - 100) for x in all_m) + 0.5
    chart.x_axis.scaling.min, chart.x_axis.scaling.max = round(100 - span_r, 1), round(100 + span_r, 1)
    chart.y_axis.scaling.min, chart.y_axis.scaling.max = round(100 - span_m, 1), round(100 + span_m, 1)
    chart.x_axis.crossesAt = 100
    chart.y_axis.crossesAt = 100
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.height, chart.width = 16, 26
    for code, a, b in ranges:
        xs = Reference(wt, min_col=3, min_row=a, max_row=b)
        ys = Reference(wt, min_col=4, min_row=a, max_row=b)
        s = Series(ys, xs, title=code + (" ★" if sector_rotation[code].get("future_star") else ""))
        s.marker = Marker(symbol="circle", size=5)
        s.smooth = False
        chart.series.append(s)
    wc.add_chart(chart, "A4")

    # --- Yontem ---
    wm = wb.create_sheet("Yöntem")
    notes = [
        "RS = 100 × Sektör / XU100 (haftalık Cuma kapanışı)",
        "RS-Ratio = 100 + (RS − 26 haftalık ortalama) / 26 haftalık std",
        "RS-Momentum = 100 + z(RS-Ratio'nun 4 haftalık değişimi), 26 hafta pencere",
        "Kadran: Lider (R>100, M>100), İyileşen (R<100, M>100), Zayıflayan (R>100, M<100), Geride (R<100, M<100)",
        "★ Erken ivme: son 3 hafta içinde Geride → İyileşen geçişi",
        "Backtest (2013–2026, 20 bps maliyet): İyileşen/erken ivme sepetleri eşit ağırlıklı sektör sepetinin yılda ~%5–8 gerisinde; 12-1 sektör momentumu ~%9–13 önünde ama istatistiksel eşiğin altında.",
        "Kaynak: BIST sektör endeksleri (borsapy / TradingView). JdK RRG'nin açık kaynak yaklaşımıdır.",
        "Sınır: pencere uzunlukları kalibre edilmedi; skora dahil değil; geçmiş performans geleceği garanti etmez.",
    ]
    wm["A1"] = "Yöntem ve Sınırlar"
    wm["A1"].font = Font(bold=True, size=12)
    for i, n in enumerate(notes, 3):
        wm.cell(row=i, column=1, value=f"• {n}")
    wm.column_dimensions["A"].width = 120
    return wb


def write_sector_rotation_excel(path: Path, as_of_date: str, sector_rotation: dict[str, dict],
                                candidates: list[dict] | None = None) -> Path | None:
    if not sector_rotation:
        return None
    wb = build_sector_rotation_workbook(as_of_date, sector_rotation, candidates)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
